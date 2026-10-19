"""Reports built with Pandas + NumPy, exportable as JSON, CSV, Excel (.xlsx) and PDF.

How every report is made (the same 4 steps - easy to explain in viva):
   1. QUERY    pull the raw rows from the database (SQLAlchemy)
   2. FRAME    put them in a pandas DataFrame
   3. SHAPE    group / count / sum / pivot with pandas; NumPy does the maths (rates, shares)
   4. PACKAGE  turn the result into one plain structure:

        {"report_type", "title", "period": {"from", "to"}, "generated_at",
         "summary": [{"label", "value"}, ...],                     <- headline numbers
         "tables":  [{"name", "title", "columns": [...], "rows": [[...], ...]}, ...]}

   That one structure is then rendered as JSON / CSV / Excel / PDF, so every format shows
   exactly the same numbers.

Dates are the HOSPITAL's local dates (utils/timeutils.py); the database holds UTC.
"""
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal

import numpy as np
import pandas as pd
from flask import current_app

from app.extensions import db
from app.models.appointment import Appointment
from app.models.bed import Admission
from app.models.billing import Bill, BillItem, Payment
from app.models.doctor import Department, Doctor
from app.models.lab import LabOrder
from app.models.patient import Patient
from app.models.pharmacy import Medicine, MedicineBatch, Prescription, PrescriptionItem
from app.models.triage import TriageRecord
from app.models.user import User
from app.services import pharmacy_service
from app.utils.timeutils import (
    date_range, local_today, month_end, month_start, range_bounds_utc, to_local_date, utc_now_naive,
)

STATUSES = ["scheduled", "completed", "cancelled", "no_show", "waitlisted"]
MAX_RANGE_DAYS = 366

# report type -> (title, parameters it accepts, permission needed)
REPORTS = {
    "daily": ("Daily Report", ["date"], "reports:all"),
    "weekly": ("Weekly Report", ["date"], "reports:all"),
    "monthly": ("Monthly Report", ["month"], "reports:all"),
    "department": ("Department Report", ["from", "to"], "reports:all"),
    "doctor-performance": ("Doctor Performance Report", ["from", "to"], "reports:all"),
    "revenue": ("Revenue Report", ["from", "to"], "reports:revenue"),
    "pharmacy": ("Pharmacy Report", ["from", "to"], "reports:pharmacy"),
}


# ============================================================================
# period parsing
# ============================================================================

def _parse_day(value, name, errors):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        errors[name] = f"{name} must be in format YYYY-MM-DD"


def resolve_period(report_type, args):
    """-> (first_day, last_day, errors). Daily = 1 day, weekly = Monday..Sunday, monthly = whole month,
    others = from..to (default: the last 30 days)."""
    errors, today = {}, local_today()
    if report_type == "daily":
        day = _parse_day(args["date"], "date", errors) if args.get("date") else today
        return day, day, errors
    if report_type == "weekly":
        day = _parse_day(args["date"], "date", errors) if args.get("date") else today
        if errors:
            return None, None, errors
        monday = day - timedelta(days=day.weekday())
        return monday, monday + timedelta(days=6), errors
    if report_type == "monthly":
        raw = args.get("month")
        if raw:
            if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", raw):
                return None, None, {"month": "month must be in format YYYY-MM"}
            day = date(int(raw[:4]), int(raw[5:]), 1)
        else:
            day = month_start(today)
        return day, month_end(day), errors
    first = _parse_day(args["from"], "from", errors) if args.get("from") else today - timedelta(days=29)
    last = _parse_day(args["to"], "to", errors) if args.get("to") else today
    if errors:
        return None, None, errors
    if first > last:
        errors["from"] = "from must not be after to"
    elif (last - first).days + 1 > MAX_RANGE_DAYS:
        errors["to"] = f"the period can cover at most {MAX_RANGE_DAYS} days"
    return first, last, errors


# ============================================================================
# small pandas helpers
# ============================================================================

def _py(v):
    """numpy / pandas / Decimal values -> plain Python (so JSON, Excel and PDF all accept them)."""
    if isinstance(v, np.generic):
        return v.item()
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (pd.Timestamp, datetime, date)):
        return v.isoformat()
    return v


def _table(name, title, df):
    df = df.copy()
    for col in df.columns:
        if pd.api.types.is_float_dtype(df[col]):
            df[col] = df[col].round(2)
    df = df.astype(object).where(pd.notna(df), None)
    return {"name": name, "title": title, "columns": [str(c) for c in df.columns],
            "rows": [[_py(v) for v in row] for row in df.values.tolist()]}


def _per_day(dates, days):
    """How many items fell on each day -> Series indexed by every day (zero-filled)."""
    counts = pd.Series(dates, dtype="object").value_counts() if dates else pd.Series(dtype="int64")
    return counts.reindex(days, fill_value=0).astype(int)


def _sum_per_day(dates, amounts, days):
    if not dates:
        return pd.Series(0.0, index=days)
    return pd.DataFrame({"d": dates, "v": amounts}).groupby("d")["v"].sum().reindex(days, fill_value=0.0).astype(float)


def _frame(rows, columns):
    return pd.DataFrame([tuple(r) for r in rows], columns=columns)


def _status_matrix(df, key):
    """Rows = value of `key`, columns = appointment statuses, cells = counts."""
    if df.empty:
        return pd.DataFrame(0, index=pd.Index([], name=key), columns=STATUSES)
    return pd.crosstab(df[key], df["status"]).reindex(columns=STATUSES, fill_value=0)


def _rate(part, whole):
    """percentage, 0 when whole is 0 (NumPy avoids a divide-by-zero)."""
    part, whole = np.asarray(part, dtype=float), np.asarray(whole, dtype=float)
    return np.round(np.where(whole > 0, part / np.where(whole > 0, whole, 1) * 100, 0.0), 1)


def _packaged(report_type, first, last, summary, tables):
    return {"report_type": report_type, "title": f"{REPORTS[report_type][0]}", "period": {"from": first.isoformat(), "to": last.isoformat()},
            "generated_at": utc_now_naive().isoformat(timespec="seconds"),
            "summary": [{"label": k, "value": _py(v)} for k, v in summary], "tables": tables}


# ============================================================================
# 1-3. DAILY / WEEKLY / MONTHLY  (hospital activity per day)
# ============================================================================

def activity_report(report_type, first, last):
    start, end = range_bounds_utc(first, last)
    days = date_range(first, last)

    def local_dates(column, extra_filter=None):
        q = db.session.query(column).filter(column >= start, column < end)
        return [to_local_date(r[0]) for r in (q.filter(extra_filter) if extra_filter is not None else q).all()]

    appts = _frame(db.session.query(Appointment.appointment_date, Appointment.status, Department.name)
                   .join(Department, Department.id == Appointment.department_id)
                   .filter(Appointment.appointment_date >= first, Appointment.appointment_date <= last).all(),
                   ["date", "status", "department"])
    booked = appts[appts["status"].isin(["scheduled", "completed", "no_show"])]
    pay_rows = db.session.query(Payment.paid_at, Payment.amount, Payment.method).filter(Payment.paid_at >= start, Payment.paid_at < end).all()
    pay_dates = [to_local_date(r[0]) for r in pay_rows]

    per_day = pd.DataFrame({
        "New patients": _per_day(local_dates(Patient.created_at), days),
        "Appointments": _per_day(list(booked["date"]), days),
        "Admissions": _per_day(local_dates(Admission.admission_date), days),
        "Discharges": _per_day(local_dates(Admission.discharge_date, Admission.discharge_date.isnot(None)), days),
        "Emergency cases": _per_day(local_dates(TriageRecord.recorded_at), days),
        "Lab tests ordered": _per_day(local_dates(LabOrder.ordered_at), days),
        "Bills generated": _per_day(local_dates(Bill.created_at), days),
        "Revenue collected (Rs.)": _sum_per_day(pay_dates, [float(r[1]) for r in pay_rows], days),
    }, index=days)
    totals = per_day.sum()
    count = lambda status: int((appts["status"] == status).sum())

    activity = per_day.drop(columns="Revenue collected (Rs.)").sum(axis=1)
    busiest = "-" if activity.max() == 0 else f"{activity.idxmax().isoformat()} ({int(activity.max())} activities)"
    summary = [("Period", f"{first.isoformat()} to {last.isoformat()}")] + [
        (col, float(totals[col]) if "Rs." in col else int(totals[col])) for col in per_day.columns] + [
        ("Appointments completed", count("completed")), ("Appointments cancelled", count("cancelled")),
        ("No-shows", count("no_show")), ("Waiting list", count("waitlisted")), ("Busiest day", busiest)]

    per_day.insert(0, "Date", [d.isoformat() for d in days])
    by_dept = _status_matrix(appts, "department")
    by_dept["Total"] = by_dept.sum(axis=1)
    by_dept = by_dept.rename(columns=lambda c: c.replace("_", "-").title()).reset_index().rename(columns={"department": "Department"})
    pay_df = _frame([(r[2], float(r[1])) for r in pay_rows], ["Method", "Amount"])
    by_method = (pay_df.groupby("Method").agg(Payments=("Amount", "size"), Amount=("Amount", "sum")).reset_index()
                 if not pay_df.empty else pd.DataFrame(columns=["Method", "Payments", "Amount"]))
    by_method = by_method.rename(columns={"Amount": "Amount (Rs.)"})

    return _packaged(report_type, first, last, summary, [
        _table("daily_breakdown", "Activity per day", per_day),
        _table("appointments_by_department", "Appointments by department", by_dept),
        _table("payments_by_method", "Payments by method", by_method)])


# ============================================================================
# 4. DEPARTMENT REPORT
# ============================================================================

def department_report(first, last):
    start, end = range_bounds_utc(first, last)
    depts = _frame([(d.id, d.name) for d in Department.query.order_by(Department.name)], ["id", "Department"]).set_index("id")
    appts = _frame(db.session.query(Appointment.department_id, Appointment.status, Appointment.patient_id)
                   .filter(Appointment.appointment_date >= first, Appointment.appointment_date <= last).all(),
                   ["department_id", "status", "patient_id"])
    matrix = _status_matrix(appts, "department_id")
    seen = appts[appts["status"].isin(["scheduled", "completed", "no_show"])]
    unique_patients = seen.groupby("department_id")["patient_id"].nunique() if not seen.empty else pd.Series(dtype="int64")
    doctors = _frame(db.session.query(Doctor.department_id, Doctor.id).filter(Doctor.is_active.is_(True)).all(), ["department_id", "id"])
    admissions = _frame(db.session.query(Doctor.department_id, Admission.id).join(Doctor, Doctor.id == Admission.doctor_id)
                        .filter(Admission.admission_date >= start, Admission.admission_date < end).all(), ["department_id", "id"])

    df = depts.join(matrix).fillna(0)
    df["Doctors"] = doctors.groupby("department_id").size().reindex(df.index, fill_value=0) if not doctors.empty else 0
    df["Unique patients"] = unique_patients.reindex(df.index, fill_value=0)
    df["Admissions"] = admissions.groupby("department_id").size().reindex(df.index, fill_value=0) if not admissions.empty else 0
    df["Booked"] = df[["scheduled", "completed", "cancelled", "no_show"]].sum(axis=1)
    df["Completion %"] = _rate(df["completed"], df["Booked"])
    out = df.rename(columns={"completed": "Completed", "cancelled": "Cancelled", "no_show": "No-show", "scheduled": "Scheduled", "waitlisted": "Waitlisted"})
    out = out[["Department", "Doctors", "Booked", "Scheduled", "Completed", "Cancelled", "No-show", "Waitlisted",
               "Completion %", "Unique patients", "Admissions"]].sort_values(["Booked", "Department"], ascending=[False, True])
    out = out.astype({c: int for c in out.columns if c not in ("Department", "Completion %")})

    busiest = out.iloc[0]["Department"] if len(out) and out.iloc[0]["Booked"] > 0 else "-"
    summary = [("Period", f"{first.isoformat()} to {last.isoformat()}"), ("Departments", len(out)),
               ("Appointments booked", int(out["Booked"].sum())), ("Appointments completed", int(out["Completed"].sum())),
               ("Admissions", int(out["Admissions"].sum())), ("Busiest department", busiest)]
    return _packaged("department", first, last, summary, [_table("departments", "Per-department activity", out)])


# ============================================================================
# 5. DOCTOR PERFORMANCE REPORT
# ============================================================================

def doctor_performance_report(first, last):
    start, end = range_bounds_utc(first, last)
    docs = _frame([(d.id, d.user.full_name, d.department.name, d.specialization) for d in Doctor.query.filter_by(is_active=True).join(User).order_by(User.full_name)],
                  ["id", "Doctor", "Department", "Specialization"]).set_index("id")
    appts = _frame(db.session.query(Appointment.doctor_id, Appointment.status, Appointment.patient_id)
                   .filter(Appointment.appointment_date >= first, Appointment.appointment_date <= last).all(),
                   ["doctor_id", "status", "patient_id"])
    matrix = _status_matrix(appts, "doctor_id")
    seen = appts[appts["status"].isin(["scheduled", "completed", "no_show"])]
    unique_patients = seen.groupby("doctor_id")["patient_id"].nunique() if not seen.empty else pd.Series(dtype="int64")

    def counted(model, doctor_col, time_col):
        rows = _frame(db.session.query(doctor_col, model.id).filter(time_col >= start, time_col < end).all(), ["doctor_id", "id"])
        return rows.groupby("doctor_id").size() if not rows.empty else pd.Series(dtype="int64")

    df = docs.join(matrix).fillna(0)
    df["Unique patients"] = unique_patients.reindex(df.index, fill_value=0)
    df["Lab tests ordered"] = counted(LabOrder, LabOrder.doctor_id, LabOrder.ordered_at).reindex(df.index, fill_value=0)
    df["Prescriptions"] = counted(Prescription, Prescription.doctor_id, Prescription.created_at).reindex(df.index, fill_value=0)
    df["Admissions"] = counted(Admission, Admission.doctor_id, Admission.admission_date).reindex(df.index, fill_value=0)
    df["Booked"] = df[["scheduled", "completed", "cancelled", "no_show"]].sum(axis=1)       # waiting list is not a booking yet
    df["Completion %"] = _rate(df["completed"], df["Booked"])
    out = df.rename(columns={"completed": "Completed", "cancelled": "Cancelled", "no_show": "No-show"})
    out = out[["Doctor", "Department", "Specialization", "Booked", "Completed", "Cancelled", "No-show", "Completion %",
               "Unique patients", "Lab tests ordered", "Prescriptions", "Admissions"]]
    out = out.sort_values(["Completed", "Doctor"], ascending=[False, True])
    out = out.astype({c: int for c in out.columns if c not in ("Doctor", "Department", "Specialization", "Completion %")})

    top = out.iloc[0]["Doctor"] if len(out) and out.iloc[0]["Completed"] > 0 else "-"
    summary = [("Period", f"{first.isoformat()} to {last.isoformat()}"), ("Doctors", len(out)),
               ("Appointments completed", int(out["Completed"].sum())), ("Prescriptions written", int(out["Prescriptions"].sum())),
               ("Lab tests ordered", int(out["Lab tests ordered"].sum())), ("Most appointments completed", top)]
    return _packaged("doctor-performance", first, last, summary, [_table("doctors", "Per-doctor activity", out)])


# ============================================================================
# 6. REVENUE REPORT
# ============================================================================

def revenue_report(first, last):
    start, end = range_bounds_utc(first, last)
    days = date_range(first, last)
    bills = _frame([(b.id, float(b.subtotal), float(b.discount_amount), float(b.tax_amount), float(b.total_amount), float(b.outstanding))
                    for b in Bill.query.filter(Bill.created_at >= start, Bill.created_at < end).all()],
                   ["id", "subtotal", "discount", "tax", "total", "outstanding"])
    items = _frame(db.session.query(BillItem.category, BillItem.amount).join(Bill, Bill.id == BillItem.bill_id)
                   .filter(Bill.created_at >= start, Bill.created_at < end).all(), ["Category", "Amount"])
    items["Amount"] = items["Amount"].astype(float)
    pays = _frame(db.session.query(Payment.paid_at, Payment.method, Payment.amount).filter(Payment.paid_at >= start, Payment.paid_at < end).all(),
                  ["paid_at", "Method", "Amount"])
    pays["Amount"] = pays["Amount"].astype(float)

    by_category = (items.groupby("Category").agg(Lines=("Amount", "size"), Amount=("Amount", "sum")).reset_index()
                   if not items.empty else pd.DataFrame(columns=["Category", "Lines", "Amount"]))
    by_category["Share %"] = _rate(by_category["Amount"], by_category["Amount"].sum()) if len(by_category) else []
    by_category = by_category.sort_values("Amount", ascending=False).rename(columns={"Amount": "Amount (Rs.)"})

    by_method = (pays.groupby("Method").agg(Payments=("Amount", "size"), Amount=("Amount", "sum")).reset_index()
                 if not pays.empty else pd.DataFrame(columns=["Method", "Payments", "Amount"]))
    by_method["Share %"] = _rate(by_method["Amount"], by_method["Amount"].sum()) if len(by_method) else []
    by_method = by_method.sort_values("Amount", ascending=False).rename(columns={"Amount": "Amount (Rs.)"})

    pay_days = [to_local_date(ts) for ts in pays["paid_at"]]
    daily = pd.DataFrame({"Date": [d.isoformat() for d in days], "Payments": _per_day(pay_days, days).values,
                          "Collected (Rs.)": _sum_per_day(pay_days, list(pays["Amount"]), days).values})

    summary = [("Period", f"{first.isoformat()} to {last.isoformat()}"), ("Bills generated", len(bills)),
               ("Gross billed (Rs.)", float(bills["subtotal"].sum())), ("Discounts given (Rs.)", float(bills["discount"].sum())),
               ("Tax charged (Rs.)", float(bills["tax"].sum())), ("Net billed (Rs.)", float(bills["total"].sum())),
               ("Collected in period (Rs.)", float(pays["Amount"].sum())),
               ("Still outstanding on these bills (Rs.)", float(bills["outstanding"].sum()))]
    return _packaged("revenue", first, last, summary, [
        _table("by_category", "Billing by category", by_category), _table("by_payment_method", "Collections by payment method", by_method),
        _table("daily_collections", "Collections per day", daily)])


# ============================================================================
# 7. PHARMACY REPORT
# ============================================================================

def pharmacy_report(first, last):
    start, end = range_bounds_utc(first, last)
    today = local_today()
    limit = today + timedelta(days=current_app.config["EXPIRY_ALERT_DAYS"])
    usable = pharmacy_service.stock_totals()

    batches = _frame([(b.medicine_id, b.medicine.name, b.batch_number, b.expiry_date, b.available_quantity)
                      for b in MedicineBatch.query.filter(MedicineBatch.available_quantity > 0).order_by(MedicineBatch.expiry_date).all()],
                     ["medicine_id", "Medicine", "Batch", "expiry", "Units"])
    expired_units = (batches[batches["expiry"] < today].groupby("medicine_id")["Units"].sum() if not batches.empty else pd.Series(dtype="int64"))

    meds = _frame([(m.id, m.name, m.unit, float(m.unit_price), m.reorder_level) for m in Medicine.query.filter_by(is_active=True).order_by(Medicine.name)],
                  ["id", "Medicine", "Unit", "Unit price (Rs.)", "Reorder level"]).set_index("id")
    meds["Usable stock"] = pd.Series(usable).reindex(meds.index, fill_value=0).astype(int)
    meds["Expired stock"] = expired_units.reindex(meds.index, fill_value=0).astype(int)
    meds["Status"] = np.select([meds["Usable stock"] == 0, meds["Usable stock"] < meds["Reorder level"]], ["OUT OF STOCK", "LOW"], default="OK")
    meds["Stock value (Rs.)"] = meds["Usable stock"] * meds["Unit price (Rs.)"]
    stock = meds[["Medicine", "Unit", "Usable stock", "Expired stock", "Reorder level", "Status", "Unit price (Rs.)", "Stock value (Rs.)"]]

    soon = batches[batches["expiry"] <= limit].copy() if not batches.empty else batches
    if not soon.empty:
        soon["Days left"] = [(e - today).days for e in soon["expiry"]]
        soon["Status"] = np.where(soon["Days left"] < 0, "EXPIRED", "Expiring soon")
        soon["Expiry date"] = [e.isoformat() for e in soon["expiry"]]
        soon = soon[["Medicine", "Batch", "Expiry date", "Units", "Days left", "Status"]]
    else:
        soon = pd.DataFrame(columns=["Medicine", "Batch", "Expiry date", "Units", "Days left", "Status"])

    used = _frame(db.session.query(Medicine.name, PrescriptionItem.quantity_issued, Prescription.id)
                  .select_from(PrescriptionItem).join(Prescription, Prescription.id == PrescriptionItem.prescription_id)
                  .join(Medicine, Medicine.id == PrescriptionItem.medicine_id)
                  .filter(Prescription.created_at >= start, Prescription.created_at < end, PrescriptionItem.quantity_issued > 0).all(),
                  ["Medicine", "issued", "prescription"])
    consumption = (used.groupby("Medicine").agg(**{"Units issued": ("issued", "sum"), "Prescriptions": ("prescription", "nunique")})
                   .reset_index().sort_values(["Units issued", "Medicine"], ascending=[False, True])
                   if not used.empty else pd.DataFrame(columns=["Medicine", "Units issued", "Prescriptions"]))

    summary = [("Period (consumption)", f"{first.isoformat()} to {last.isoformat()}"), ("Medicines in catalogue", len(stock)),
               ("Out of stock", int((stock["Status"] == "OUT OF STOCK").sum())), ("Low stock", int((stock["Status"] == "LOW").sum())),
               ("Batches expiring soon", int((soon["Status"] == "Expiring soon").sum())),
               ("Expired batches still in stock", int((soon["Status"] == "EXPIRED").sum())),
               ("Units issued in period", int(consumption["Units issued"].sum()) if len(consumption) else 0),
               ("Stock value (Rs.)", float(stock["Stock value (Rs.)"].sum()))]
    return _packaged("pharmacy", first, last, summary, [
        _table("stock", "Stock on hand", stock), _table("expiring_batches", f"Batches expiring within {current_app.config['EXPIRY_ALERT_DAYS']} days (or already expired)", soon),
        _table("consumption", "Consumption in period", consumption)])


BUILDERS = {"department": department_report, "doctor-performance": doctor_performance_report,
            "revenue": revenue_report, "pharmacy": pharmacy_report}


def build_report(report_type, first, last):
    if report_type in ("daily", "weekly", "monthly"):
        return activity_report(report_type, first, last)
    return BUILDERS[report_type](first, last)


# ============================================================================
# EXPORTS
# ============================================================================

def _safe_cell(value):
    """Spreadsheet formula injection guard: text starting with = + @ (or a control char) would be
    EXECUTED by Excel when the file is opened, so we neutralise it with a leading apostrophe."""
    if isinstance(value, str) and (value[:1] in ("=", "+", "@", "\t", "\r") or (value[:1] == "-" and len(value) > 1)):
        return "'" + value
    return value


def _safe_frame(table):
    return pd.DataFrame(table["rows"], columns=table["columns"]).map(_safe_cell)


def table_names(report):
    return [t["name"] for t in report["tables"]]


def export_csv(report, table=None):
    """CSV holds ONE table. Default: the first one; choose another with ?table=<name>.
    Written as UTF-8 with a BOM so Excel shows non-English names correctly."""
    chosen = next(t for t in report["tables"] if table is None or t["name"] == table)
    return _safe_frame(chosen).to_csv(index=False).encode("utf-8-sig")


def _sheet_name(name, used):
    base = re.sub(r"[\[\]:*?/\\]", "", name)[:31] or "Sheet"
    candidate, n = base, 2
    while candidate in used:
        candidate = f"{base[:28]}_{n}"
        n += 1
    used.add(candidate)
    return candidate


def export_xlsx(report):
    from openpyxl.styles import Font
    buffer, used = io.BytesIO(), set()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary = pd.DataFrame([[s["label"], s["value"]] for s in report["summary"]], columns=["Metric", "Value"]).map(_safe_cell)
        name = _sheet_name("Summary", used)
        summary.to_excel(writer, sheet_name=name, index=False, startrow=3)
        ws = writer.sheets[name]
        ws["A1"] = f'{current_app.config["HOSPITAL_NAME"]} - {report["title"]}'
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f'Period: {report["period"]["from"]} to {report["period"]["to"]}   |   Generated: {report["generated_at"]} UTC'
        for table in report["tables"]:
            sheet = _sheet_name(table["title"], used)
            _safe_frame(table).to_excel(writer, sheet_name=sheet, index=False)
            writer.sheets[sheet].freeze_panes = "A2"
        for ws in writer.book.worksheets:
            header_row = 4 if ws.title == "Summary" else 1
            for cell in ws[header_row]:
                cell.font = Font(bold=True)
            for column in ws.columns:
                longest = max((len(str(c.value)) for c in column if c.value is not None and c.row >= header_row), default=8)
                ws.column_dimensions[column[0].column_letter].width = min(max(longest + 2, 10), 60)
    return buffer.getvalue()
