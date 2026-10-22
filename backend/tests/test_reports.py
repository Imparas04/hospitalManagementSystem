import io
from datetime import timedelta

import pandas as pd
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.utils.timeutils import local_today

ALL = ["daily", "weekly", "monthly", "department", "doctor-performance", "revenue", "pharmacy"]
ROLES = ["admin", "receptionist", "doctor", "nurse", "lab_technician", "pharmacist", "billing_staff"]
ALLOWED = {"revenue": {"admin", "billing_staff"}, "pharmacy": {"admin", "pharmacist"}}


def _get(client, auth_header, role, report, **params):
    q = "&".join(f"{k}={v}" for k, v in params.items())
    return client.get(f"/api/reports/{report}?{q}", headers=auth_header(client, role))


def _json(client, auth_header, report, role="admin", **params):
    return _get(client, auth_header, role, report, **params).get_json()["data"]


def _summary(report):
    return {s["label"]: s["value"] for s in report["summary"]}


def _table(report, name):
    t = next(t for t in report["tables"] if t["name"] == name)
    return [dict(zip(t["columns"], row)) for row in t["rows"]]


def _week():
    today = local_today()
    return today - timedelta(days=today.weekday())


# ------------------------------------------------------------------- access --

def test_report_list_depends_on_role(client, auth_header):
    def types(role):
        return [r["type"] for r in client.get("/api/reports", headers=auth_header(client, role)).get_json()["data"]]
    assert types("admin") == ALL
    assert types("billing_staff") == ["revenue"] and types("pharmacist") == ["pharmacy"]
    for role in ("receptionist", "doctor", "nurse", "lab_technician"):
        assert types(role) == []


@pytest.mark.parametrize("report,role", [(r, role) for r in ALL for role in ROLES])
def test_report_access_matrix(client, auth_header, report, role):
    r = _get(client, auth_header, role, report)
    assert r.status_code == (200 if role in ALLOWED.get(report, {"admin"}) else 403)


def test_unknown_report_and_no_token(client, auth_header):
    assert _get(client, auth_header, "admin", "bogus").status_code == 404
    assert client.get("/api/reports/daily").status_code == 401
    assert client.get("/api/reports").status_code == 401


@pytest.mark.parametrize("report", ALL)
@pytest.mark.parametrize("fmt", ["json", "csv", "xlsx", "pdf"])
def test_every_report_works_on_an_empty_hospital(client, auth_header, report, fmt):
    """The nastiest case for pandas: no rows at all."""
    r = _get(client, auth_header, "admin", report, format=fmt)
    assert r.status_code == 200 and len(r.data) > 0
    if fmt == "csv":
        assert b"," in r.data.splitlines()[0]                          # at least the header row
    elif fmt == "json":
        assert r.get_json()["data"]["tables"] and r.get_json()["data"]["summary"]
    elif fmt == "pdf":
        assert r.data.startswith(b"%PDF")
    else:
        assert load_workbook(io.BytesIO(r.data)).sheetnames[0] == "Summary"


# --------------------------------------------------------------- the numbers --

def test_daily_report_counts_the_day(client, auth_header, scenario, patient_id):
    scenario()
    h = auth_header(client, "billing_staff")
    bill = client.post("/api/bills", headers=h, json={"patient_id": patient_id}).get_json()["data"]
    client.post(f"/api/bills/{bill['id']}/payments", headers=h, json={"amount": 1000, "method": "cash"})
    rep = _json(client, auth_header, "daily")
    s = _summary(rep)
    assert rep["period"] == {"from": local_today().isoformat(), "to": local_today().isoformat()}
    assert s["New patients"] == 1 and s["Appointments"] == 1 and s["Admissions"] == 1 and s["Discharges"] == 1
    assert s["Emergency cases"] == 0 and s["Lab tests ordered"] == 1 and s["Bills generated"] == 1
    assert s["Revenue collected (Rs.)"] == 1000.0 and s["Appointments completed"] == 1
    assert [t["name"] for t in rep["tables"]] == ["daily_breakdown", "appointments_by_department", "payments_by_method"]
    assert _table(rep, "payments_by_method") == [{"Method": "cash", "Payments": 1, "Amount (Rs.)": 1000.0}]
    dept = _table(rep, "appointments_by_department")[0]
    assert dept["Department"] == "Cardiology" and dept["Completed"] == 1 and dept["Total"] == 1


def test_daily_report_for_another_date_is_empty(client, auth_header, scenario):
    scenario()
    s = _summary(_json(client, auth_header, "daily", date="2020-01-01"))
    assert s["New patients"] == 0 and s["Appointments"] == 0 and s["Busiest day"] == "-"


def test_weekly_report_runs_monday_to_sunday(client, auth_header, scenario):
    scenario()
    monday = _week()
    for day in (monday, monday + timedelta(days=3), monday + timedelta(days=6)):          # Mon, Thu, Sun -> same week
        rep = _json(client, auth_header, "weekly", date=day.isoformat())
        assert rep["period"] == {"from": monday.isoformat(), "to": (monday + timedelta(days=6)).isoformat()}
    rows = _table(rep, "daily_breakdown")
    assert len(rows) == 7 and rows[0]["Date"] == monday.isoformat()
    assert _summary(rep)["New patients"] == 1 and sum(r["Appointments"] for r in rows) == 1


def test_monthly_report(client, auth_header, scenario):
    scenario()
    today = local_today()
    rep = _json(client, auth_header, "monthly", month=f"{today.year}-{today.month:02d}")
    assert rep["period"]["from"] == today.replace(day=1).isoformat()
    assert len(_table(rep, "daily_breakdown")) == (pd.Period(today, "M").days_in_month)
    assert _summary(rep)["Admissions"] == 1
    default = _json(client, auth_header, "monthly")                    # no month -> this month
    assert default["period"] == rep["period"]


def test_department_report(client, auth_header, scenario):
    scenario()
    client.post("/api/departments", headers=auth_header(client, "admin"), json={"name": "Orthopedics"})
    rep = _json(client, auth_header, "department")
    rows = {r["Department"]: r for r in _table(rep, "departments")}
    c = rows["Cardiology"]
    assert (c["Doctors"], c["Booked"], c["Completed"], c["Completion %"], c["Unique patients"], c["Admissions"]) == (1, 1, 1, 100.0, 1, 1)
    o = rows["Orthopedics"]                                            # nothing happened: zeros, and no divide-by-zero
    assert (o["Booked"], o["Completed"], o["Completion %"], o["Unique patients"]) == (0, 0, 0.0, 0)
    assert list(rows)[0] == "Cardiology"                               # busiest first
    s = _summary(rep)
    assert s["Appointments booked"] == 1 and s["Busiest department"] == "Cardiology" and s["Departments"] == 2


def test_doctor_performance_report(client, auth_header, scenario, doctor_setup):
    scenario()
    admin = auth_header(client, "admin")
    client.post("/api/users", headers=admin, json={"username": "doc2", "email": "doc2@test.com", "password": "Secret123",
                                                    "full_name": "Asha Rao", "role": "doctor"})
    uid = client.get("/api/users?q=doc2", headers=admin).get_json()["data"]["items"][0]["id"]
    client.post("/api/doctors", headers=admin, json={"user_id": uid, "department_id": doctor_setup["department"]["id"], "specialization": "ENT"})
    rows = _table(_json(client, auth_header, "doctor-performance"), "doctors")
    assert [r["Doctor"] for r in rows] == ["Test doctor", "Asha Rao"]          # most completed first
    first, idle = rows
    assert (first["Booked"], first["Completed"], first["Completion %"], first["Unique patients"]) == (1, 1, 100.0, 1)
    assert (first["Lab tests ordered"], first["Prescriptions"], first["Admissions"]) == (1, 1, 1)
    assert all(idle[k] == 0 for k in ("Booked", "Completed", "Completion %", "Lab tests ordered", "Prescriptions", "Admissions"))


def test_revenue_report(client, auth_header, scenario, patient_id):
    scenario()
    h = auth_header(client, "billing_staff")
    bill = client.post("/api/bills", headers=h, json={"patient_id": patient_id, "discount_percent": 10, "tax_percent": 5}).get_json()["data"]
    client.post(f"/api/bills/{bill['id']}/payments", headers=h, json={"amount": 1000, "method": "cash"})
    client.post(f"/api/bills/{bill['id']}/payments", headers=h, json={"amount": 500, "method": "upi"})
    rep = _json(client, auth_header, "revenue", role="billing_staff")
    s = _summary(rep)
    assert s["Bills generated"] == 1 and s["Gross billed (Rs.)"] == 2400.0 and s["Discounts given (Rs.)"] == 240.0
    assert s["Tax charged (Rs.)"] == 108.0 and s["Net billed (Rs.)"] == 2268.0
    assert s["Collected in period (Rs.)"] == 1500.0 and s["Still outstanding on these bills (Rs.)"] == 768.0
    cat = {r["Category"]: r for r in _table(rep, "by_category")}
    assert cat["bed"]["Amount (Rs.)"] == 1500.0 and cat["bed"]["Share %"] == 62.5
    assert sum(r["Amount (Rs.)"] for r in cat.values()) == 2400.0
    method = {r["Method"]: r for r in _table(rep, "by_payment_method")}
    assert method["cash"]["Share %"] == 66.7 and method["upi"]["Share %"] == 33.3
    daily = _table(rep, "daily_collections")
    assert len(daily) == 30 and daily[-1]["Collected (Rs.)"] == 1500.0 and sum(r["Payments"] for r in daily) == 2


def test_pharmacy_report(client, auth_header, scenario):
    from app.extensions import db
    from app.models.pharmacy import MedicineBatch
    scenario()
    admin = auth_header(client, "admin")
    today = local_today()

    def med(name, reorder, batches):
        m = client.post("/api/pharmacy/medicines", headers=admin, json={"name": name, "unit": "tablet", "reorder_level": reorder, "unit_price": 2}).get_json()["data"]
        for no, expiry, qty in batches:
            client.post(f"/api/pharmacy/medicines/{m['id']}/batches", headers=admin, json={"batch_number": no, "expiry_date": expiry, "purchase_quantity": qty})
        return m
    med("Lowcillin", 50, [("L1", "2028-01-01", 5)])
    med("Soonex", 5, [("S1", (today + timedelta(days=10)).isoformat(), 40)])
    med("Oldmed", 5, [("O1", "2028-01-01", 30)])
    MedicineBatch.query.filter_by(batch_number="O1").first().expiry_date = today - timedelta(days=3)
    db.session.commit()

    rep = _json(client, auth_header, "pharmacy", role="pharmacist")
    stock = {r["Medicine"]: r for r in _table(rep, "stock")}
    assert stock["Paracetamol"]["Usable stock"] == 16 and stock["Paracetamol"]["Status"] == "OK"      # 20 bought - 4 issued
    assert stock["Lowcillin"]["Status"] == "LOW"
    assert stock["Oldmed"]["Status"] == "OUT OF STOCK" and stock["Oldmed"]["Expired stock"] == 30 and stock["Oldmed"]["Usable stock"] == 0
    assert stock["Soonex"]["Stock value (Rs.)"] == 80.0
    batches = {r["Batch"]: r for r in _table(rep, "expiring_batches")}
    assert batches["S1"]["Status"] == "Expiring soon" and batches["S1"]["Days left"] == 10
    assert batches["O1"]["Status"] == "EXPIRED" and batches["O1"]["Days left"] == -3 and "L1" not in batches
    assert _table(rep, "consumption") == [{"Medicine": "Paracetamol", "Units issued": 4, "Prescriptions": 1}]
    s = _summary(rep)
    assert s["Out of stock"] == 1 and s["Low stock"] == 1 and s["Batches expiring soon"] == 1 and s["Expired batches still in stock"] == 1
    assert s["Units issued in period"] == 4 and s["Medicines in catalogue"] == 4


def test_late_night_payment_lands_on_the_local_day(app, client, auth_header, patient_id):
    """01:30 IST on Oct 11 = 20:00 UTC on Oct 10: the Oct 11 report must contain it, the Oct 10 report must not."""
    from datetime import datetime
    from app.extensions import db
    from app.models.billing import Payment
    h = auth_header(client, "billing_staff")
    bill = client.post("/api/bills", headers=h, json={"patient_id": patient_id, "other_services": [{"description": "Fee", "unit_price": 300}]}).get_json()["data"]
    client.post(f"/api/bills/{bill['id']}/payments", headers=h, json={"amount": 300, "method": "cash"})
    Payment.query.first().paid_at = datetime(2026, 10, 10, 20, 0)
    db.session.commit()
    app.config["HOSPITAL_TIMEZONE"] = "Asia/Kolkata"
    assert _summary(_json(client, auth_header, "daily", date="2026-10-11"))["Revenue collected (Rs.)"] == 300.0
    assert _summary(_json(client, auth_header, "daily", date="2026-10-10"))["Revenue collected (Rs.)"] == 0.0


# ------------------------------------------------------------------ validation --

@pytest.mark.parametrize("report,params,field", [
    ("daily", {"date": "10-10-2026"}, "date"), ("weekly", {"date": "tomorrow"}, "date"), ("monthly", {"month": "2026-13"}, "month"),
    ("monthly", {"month": "2026-1"}, "month"), ("department", {"from": "2026-10-10", "to": "2026-10-01"}, "from"),
    ("revenue", {"from": "2024-01-01", "to": "2026-01-01"}, "to"), ("pharmacy", {"from": "nope"}, "from")])
def test_period_validation(client, auth_header, report, params, field):
    r = _get(client, auth_header, "admin", report, **params)
    assert r.status_code == 400 and field in r.get_json()["errors"]


def test_bad_format_and_table(client, auth_header):
    assert _get(client, auth_header, "admin", "daily", format="docx").status_code == 400
    r = _get(client, auth_header, "admin", "daily", format="csv", table="nope")
    assert r.status_code == 400 and "table" in r.get_json()["errors"]


# --------------------------------------------------------------------- exports --

def test_csv_export(client, auth_header, scenario):
    scenario()
    r = _get(client, auth_header, "admin", "department", format="csv")
    assert r.status_code == 200 and r.mimetype == "text/csv"
    assert "attachment" in r.headers["Content-Disposition"] and ".csv" in r.headers["Content-Disposition"]
    assert r.data.startswith(b"\xef\xbb\xbf")                          # UTF-8 BOM: Excel shows non-English names correctly
    df = pd.read_csv(io.BytesIO(r.data), encoding="utf-8-sig")
    assert df.loc[df["Department"] == "Cardiology", "Completed"].iloc[0] == 1
    other = pd.read_csv(io.BytesIO(_get(client, auth_header, "admin", "daily", format="csv", table="payments_by_method").data), encoding="utf-8-sig")
    assert list(other.columns) == ["Method", "Payments", "Amount (Rs.)"]


def test_xlsx_export(client, auth_header, scenario):
    scenario()
    r = _get(client, auth_header, "admin", "doctor-performance", format="xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "Per-doctor activity"]
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=5, values_only=True)}
    assert summary["Appointments completed"] == 1 and "Doctor Performance Report" in wb["Summary"]["A1"].value
    sheet = wb["Per-doctor activity"]
    header = [c.value for c in sheet[1]]
    row = dict(zip(header, [c.value for c in sheet[2]]))
    assert row["Doctor"] == "Test doctor" and row["Completed"] == 1 and sheet["A1"].font.bold


def test_pdf_export(client, auth_header, scenario):
    scenario()
    r = _get(client, auth_header, "admin", "department", format="pdf")
    assert r.status_code == 200 and r.mimetype == "application/pdf" and r.data.startswith(b"%PDF")
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(r.data)).pages)
    for expected in ("DEPARTMENT REPORT", "Cardiology", "Per-department activity", "Appointments booked", "100.00"):
        assert expected in text


def test_wide_reports_print_in_landscape(client, auth_header, scenario):
    scenario()
    page = PdfReader(io.BytesIO(_get(client, auth_header, "admin", "doctor-performance", format="pdf").data)).pages[0]
    assert float(page.mediabox.width) > float(page.mediabox.height)
    narrow = PdfReader(io.BytesIO(_get(client, auth_header, "billing_staff", "revenue", format="pdf").data)).pages[0]
    assert float(narrow.mediabox.width) < float(narrow.mediabox.height)


def test_spreadsheet_formula_injection_is_neutralised(client, auth_header):
    """A department called =HYPERLINK(...) would RUN when someone opens the Excel/CSV file."""
    evil = '=HYPERLINK("http://evil.example","click")'
    client.post("/api/departments", headers=auth_header(client, "admin"), json={"name": evil})
    assert evil in str(_json(client, auth_header, "department"))                     # JSON keeps the real value
    csv = _get(client, auth_header, "admin", "department", format="csv").data.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv and ",=HYPERLINK" not in csv and not any(line.startswith("=") for line in csv.splitlines())
    wb = load_workbook(io.BytesIO(_get(client, auth_header, "admin", "department", format="xlsx").data))
    cells = [c.value for row in wb["Per-department activity"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(v.startswith("'=HYPERLINK") for v in cells) and not any(v.startswith("=") for v in cells)


def test_exports_are_audited_but_json_is_not(client, auth_header):
    admin = auth_header(client, "admin")
    _get(client, auth_header, "admin", "daily")                                       # JSON on screen: not an export
    _get(client, auth_header, "admin", "revenue", format="xlsx")
    _get(client, auth_header, "admin", "department", format="pdf")
    logs = client.get("/api/audit-logs?action=report.export", headers=admin).get_json()["data"]["items"]
    assert sorted(l["description"] for l in logs) == ["Exported department report as pdf", "Exported revenue report as xlsx"]
    assert all(l["entity_type"] == "Report" and l["role"] == "admin" for l in logs)


def test_forbidden_export_is_not_audited(client, auth_header):
    _get(client, auth_header, "doctor", "revenue", format="xlsx")                     # 403
    assert client.get("/api/audit-logs?action=report.export", headers=auth_header(client, "admin")).get_json()["data"]["total"] == 0
